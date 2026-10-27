"""Escribe los archivos de la carpeta `datos/`."""

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter

from .ensuciar import ensuciar_ventas
from .generar import SEMILLA, generar_todo, rng_tabla

# Con BOM para que Excel muestre bien las tildes al abrir el CSV con doble clic.
CODIFICACION = "utf-8-sig"
# Fecha fija en el libro de Excel, para que regenerarlo no cambie el archivo.
FECHA_FIJA = datetime(2026, 1, 1)


def exportar(carpeta: Path, semilla: int = SEMILLA) -> list[Path]:
    """Genera todas las tablas y las escribe en `carpeta`. Devuelve las rutas escritas."""
    carpeta = Path(carpeta)
    carpeta.mkdir(parents=True, exist_ok=True)
    tablas = generar_todo(semilla)
    escritos = []

    for nombre, df in tablas.items():
        ruta = carpeta / f"{nombre}.csv"
        df.to_csv(ruta, index=False, date_format="%Y-%m-%d", encoding=CODIFICACION)
        escritos.append(ruta)

    # Archivo "como lo exporta la oficina", separado por punto y coma.
    crudo = ensuciar_ventas(tablas["ventas"], rng_tabla("crudo", semilla))
    ruta = carpeta / "ventas_crudo.csv"
    crudo.to_csv(ruta, index=False, sep=";", encoding=CODIFICACION)
    escritos.append(ruta)

    ventas = tablas["ventas"]
    hojas = {
        "Ventas": ventas[ventas["fecha"].dt.year == 2025],
        "Productos": tablas["productos"],
        "Clientes": tablas["clientes"],
    }
    ruta = carpeta / "reporte_ventas_2025.xlsx"
    escribir_excel(hojas, ruta)
    escritos.append(ruta)
    return escritos


def escribir_excel(hojas: dict[str, pd.DataFrame], ruta: Path) -> None:
    """Libro de Excel con formato de oficina: fechas dd/mm/aaaa, filtros y encabezado fijo."""
    with pd.ExcelWriter(
        ruta, engine="openpyxl", date_format="DD/MM/YYYY", datetime_format="DD/MM/YYYY"
    ) as libro:
        for nombre, df in hojas.items():
            df.to_excel(libro, sheet_name=nombre, index=False)
            hoja = libro.sheets[nombre]
            hoja.freeze_panes = "A2"
            hoja.auto_filter.ref = hoja.dimensions
            for i, columna in enumerate(df.columns, start=1):
                largo = max(len(str(columna)), df[columna].astype(str).str.len().max())
                hoja.column_dimensions[get_column_letter(i)].width = min(largo + 3, 45)
                if columna == "descuento":
                    for (celda,) in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                        celda.number_format = "0%"
                elif columna in ("precio", "costo", "precio_unitario", "total"):
                    for (celda,) in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                        celda.number_format = "#,##0.00"
        libro.book.properties.creator = "Distribuidora Andina S.A.C."
        libro.book.properties.created = FECHA_FIJA
        libro.book.properties.modified = FECHA_FIJA
    _fijar_fechas_del_zip(ruta)


def _fijar_fechas_del_zip(ruta: Path) -> None:
    """Un .xlsx es un .zip: fija la fecha de cada archivo interno para que sea reproducible."""
    original = ruta.read_bytes()
    salida = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(original)) as entrada,
        zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as nuevo,
    ):
        for info in entrada.infolist():
            datos = entrada.read(info.filename)
            if info.filename == "docProps/core.xml":
                datos = _fijar_fecha_modificacion(datos)
            fija = zipfile.ZipInfo(info.filename, date_time=FECHA_FIJA.timetuple()[:6])
            fija.compress_type = zipfile.ZIP_DEFLATED
            nuevo.writestr(fija, datos)
    ruta.write_bytes(salida.getvalue())


def _fijar_fecha_modificacion(xml: bytes) -> bytes:
    fija = FECHA_FIJA.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
    return re.sub(
        rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>" + fija + rb"\g<2>", xml
    )
